from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from dsystem.export.columns import DATE, MONEY, QUANTITY, Column
from dsystem.export.response import TRUNCATED_HEADER, attachment, content_disposition
from dsystem.export.sheet import Sheet, within_limit
from dsystem.export.xlsx import XLSX_MEDIA_TYPE, render_xlsx, workbook_bytes
from dsystem.i18n import bind_locale_dir

TASHKENT = ZoneInfo("Asia/Tashkent")
COLUMNS = (
    Column("document_code", 18),
    Column("document_date", 18, DATE),
    Column("quantity", 12, QUANTITY),
    Column("total", 16, MONEY, total=True),
)


@pytest.fixture(autouse=True)
def _locale():
    bind_locale_dir(Path(__file__).resolve().parents[1] / "dsystem" / "locale")


def _sheet(rows, **overrides) -> Sheet:
    fields = {
        "title": "Sotuvlar",
        "columns": COLUMNS,
        "rows": rows,
        "lang": "uz",
        "zone": TASHKENT,
        "generated_at": datetime(2026, 10, 1, 9, 30, tzinfo=timezone.utc),
    }
    return Sheet(**(fields | overrides))


def _read(content: bytes):
    return list(load_workbook(BytesIO(content), data_only=True).active.iter_rows(values_only=True))


def _rows():
    return [
        ("SO-1", datetime(2026, 9, 30, 20, 0, tzinfo=timezone.utc), Decimal("2.5"), Decimal("100.10")),
        ("SO-2", datetime(2026, 9, 29, 8, 0, tzinfo=timezone.utc), 1, Decimal("50.05")),
    ]


def test_layout_is_title_meta_header_rows_then_total():
    rows = _read(workbook_bytes(_sheet(_rows(), date_from=date(2026, 9, 1), date_to=date(2026, 9, 30))))

    assert rows[0][0] == "Sotuvlar · 01.09.2026 – 30.09.2026"
    assert rows[1][0] == "Yuklab olindi: 01.10.2026 14:30"
    assert rows[2][:2] == ("document_code", "document_date")
    assert [r[0] for r in rows[3:5]] == ["SO-1", "SO-2"]
    assert rows[5][0] == "Jami" and rows[5][3] == pytest.approx(150.15)


def test_dates_land_in_the_viewers_zone_and_numbers_stay_numbers():
    rows = _read(workbook_bytes(_sheet(_rows())))

    assert rows[3][1] == datetime(2026, 10, 1, 1, 0)
    assert rows[3][2] == 2.5 and isinstance(rows[3][3], float)


def test_total_is_a_subtotal_formula_so_excel_filters_recount():
    book = load_workbook(BytesIO(workbook_bytes(_sheet(_rows()))))

    assert book.active.cell(row=6, column=4).value == "=SUBTOTAL(9,D4:D5)"


def test_empty_export_still_has_header_and_zero_total():
    rows = _read(workbook_bytes(_sheet([])))

    assert len(rows) == 4 and rows[3][0] == "Jami" and rows[3][3] == 0


def test_text_that_looks_like_a_formula_is_written_as_text():
    content = workbook_bytes(_sheet([('=HYPERLINK("http://x")', None, None, Decimal("1"))]))

    assert load_workbook(BytesIO(content)).active.cell(row=4, column=1).data_type == "s"


def test_truncated_export_says_so_under_the_title():
    rows = _read(workbook_bytes(_sheet(_rows(), truncated=True, lang="ru")))

    assert rows[1][0] == "Выгружено 01.10.2026 14:30 · только первые 50 000 строк"


def test_within_limit_cuts_and_flags():
    assert within_limit([1, 2, 3], limit=2) == ([1, 2], True)
    assert within_limit([1, 2], limit=2) == ([1, 2], False)


def test_sheet_name_drops_characters_excel_rejects():
    book = load_workbook(BytesIO(workbook_bytes(_sheet([], title="A/B: [x]?" + "y" * 40))))

    title = book.active.title
    assert len(title) == 31 and title.startswith("A B") and not set("[]:*?/\\") & set(title)


async def test_render_runs_off_the_event_loop_and_accepts_a_generator():
    content = await render_xlsx(_sheet(row for row in _rows()))

    assert [r[0] for r in _read(content)[3:5]] == ["SO-1", "SO-2"]


def test_content_disposition_keeps_non_ascii_names():
    header = content_disposition("Sotuvlar-Продажи.xlsx")

    assert header.startswith('attachment; filename="Sotuvlar-.xlsx"')
    assert unquote(header.split("filename*=UTF-8''", 1)[1]) == "Sotuvlar-Продажи.xlsx"


def test_attachment_flags_truncation_only_when_cut():
    assert TRUNCATED_HEADER not in attachment(b"x", "a.xlsx", XLSX_MEDIA_TYPE).headers
    assert attachment(b"x", "a.xlsx", XLSX_MEDIA_TYPE, truncated=True).headers[TRUNCATED_HEADER] == "true"


def test_columns_widen_to_fit_large_amounts_and_cap_long_text():
    columns = (Column("partner", 20), Column("total", 16, MONEY, total=True))
    rows = [("x" * 200, Decimal("3914224450360.00"))] * 3
    book = load_workbook(BytesIO(workbook_bytes(_sheet(rows, columns=columns))))
    widths = {letter: dim.width for letter, dim in book.active.column_dimensions.items()}

    assert 60 <= widths["A"] < 62
    assert widths["B"] >= len("11,742,673,351,080.00")
